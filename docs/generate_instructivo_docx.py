#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Genera el instructivo en formato Word (.docx)."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Mm, Pt, RGBColor
from PIL import Image as PILImage

from generate_instructivo import ASSETS, ensure_screenshots
from instructivo_content import TOC, render_instructivo
from pdf_assets import draw_flujo_clinico

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "INSTRUCTIVO_HealthTiCloud_RIS.docx"

BRAND = RGBColor(46, 196, 182)
BRAND_DARK = RGBColor(33, 37, 41)
MUTED = RGBColor(108, 117, 125)
HEADER_FILL = "2EC4B6"
ROW_FILL = "F8FCFB"


def _set_cell_shading(cell, fill_hex: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill_hex)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def _set_cell_margins(cell, top=60, bottom=60, left=100, right=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)


def _add_toc_field(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run._r.append(fld_begin)

    run2 = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    run2._r.append(instr)

    run3 = paragraph.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run3._r.append(fld_sep)

    run4 = paragraph.add_run("Actualice el índice: clic derecho → Actualizar campo.")
    run4.italic = True
    run4.font.color.rgb = MUTED
    run4.font.size = Pt(10)

    run5 = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run5._r.append(fld_end)


class InstructivoDOCX:
    def __init__(self) -> None:
        self.doc = Document()
        self._section_title = ""
        self._figure_no = 0
        self._setup_page()
        self._setup_styles()

    def _setup_page(self) -> None:
        section = self.doc.sections[0]
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        self._content_width = section.page_width - section.left_margin - section.right_margin

    def _setup_styles(self) -> None:
        normal = self.doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        normal.paragraph_format.line_spacing = 1.15

        for level, size, color in ((1, 18, BRAND), (2, 13, BRAND_DARK), (3, 12, BRAND_DARK)):
            style = self.doc.styles[f"Heading {level}"]
            style.font.name = "Calibri"
            style.font.bold = True
            style.font.size = Pt(size)
            style.font.color.rgb = color
            style.paragraph_format.space_before = Pt(12 if level > 1 else 0)
            style.paragraph_format.space_after = Pt(8)
            style.paragraph_format.keep_with_next = True

    def _add_paragraph(self, text: str = "", *, bold: bool = False, italic: bool = False,
                       align=WD_ALIGN_PARAGRAPH.LEFT, color: RGBColor | None = None,
                       size: float | None = None, space_after: float = 6) -> None:
        p = self.doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        if color:
            run.font.color.rgb = color
        if size:
            run.font.size = Pt(size)

    def cover(self, banner: Path | None) -> None:
        for _ in range(4):
            self.doc.add_paragraph()

        self._add_paragraph(
            "HealthTiCloud RIS",
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            color=BRAND,
            size=28,
            space_after=10,
        )
        self._add_paragraph(
            "Manual de usuario",
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            color=BRAND_DARK,
            size=16,
            space_after=4,
        )
        self._add_paragraph(
            "Instructivo por perfiles y flujo clínico",
            align=WD_ALIGN_PARAGRAPH.CENTER,
            color=BRAND_DARK,
            size=13,
            space_after=14,
        )

        if banner and banner.exists():
            self.figure(banner, "", max_height=70, caption_below=False)

        self._add_paragraph(
            "Sistema de Información Radiológica  ·  Mayo 2026",
            italic=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            color=MUTED,
            size=10,
        )
        self.doc.add_page_break()

    def toc(self, items: list[str]) -> None:
        self.doc.add_heading("Índice", level=1)
        _add_toc_field(self.doc)
        self.doc.add_paragraph()
        self._add_paragraph(
            "Índice estático (referencia rápida):",
            bold=True,
            size=11,
            space_after=4,
        )
        for title in items:
            p = self.doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_after = Pt(2)
            p.add_run(title)

    def chapter(self, number: str, title: str) -> None:
        self._section_title = title
        self.doc.add_page_break()
        self.doc.add_heading(f"{number}. {title}", level=1)

    def h2(self, text: str) -> None:
        self.doc.add_heading(text, level=2)

    def p(self, text: str) -> None:
        self._add_paragraph(text)

    def bullets(self, items: list[str]) -> None:
        for item in items:
            p = self.doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            p.add_run(item)

    def steps(self, items: list[str]) -> None:
        for i, item in enumerate(items, 1):
            p = self.doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(3)
            p.add_run(item)

    def figure(
        self,
        path: Path,
        caption: str,
        max_height: float = 88,
        caption_below: bool = True,
    ) -> None:
        if not path.exists():
            return

        with PILImage.open(path) as img:
            w_px, h_px = img.size
        aspect = h_px / w_px if w_px else 0.56

        max_w = self._content_width
        max_h = Mm(max_height)
        img_w = max_w
        img_h = int(img_w * aspect)
        if img_h > max_h:
            img_h = max_h
            img_w = int(img_h / aspect)

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run()
        run.add_picture(str(path), width=img_w, height=img_h)

        if caption and caption_below:
            self._figure_no += 1
            cap = self.doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_after = Pt(10)
            run_cap = cap.add_run(caption)
            run_cap.italic = True
            run_cap.font.size = Pt(10)
            run_cap.font.color.rgb = MUTED

    def render_table(
        self,
        headers: list[str],
        rows: list[list[str]],
        col_widths: tuple[float, ...] | None = None,
    ) -> None:
        n = len(headers)
        table = self.doc.add_table(rows=1 + len(rows), cols=n)
        table.style = "Table Grid"
        table.autofit = False

        total_mm = 166.0
        if col_widths:
            total = sum(col_widths)
            widths = [Mm(total_mm * w / total) for w in col_widths]
        else:
            widths = [Mm(total_mm / n)] * n

        for col_idx, width in enumerate(widths):
            for row in table.rows:
                row.cells[col_idx].width = width

        hdr_cells = table.rows[0].cells
        for idx, header in enumerate(headers):
            cell = hdr_cells[idx]
            _set_cell_shading(cell, HEADER_FILL)
            _set_cell_margins(cell)
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(header)
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(10)

        for r_idx, row_data in enumerate(rows, start=1):
            row = table.rows[r_idx]
            if r_idx % 2 == 0:
                for cell in row.cells:
                    _set_cell_shading(cell, ROW_FILL)
            for c_idx, value in enumerate(row_data):
                cell = row.cells[c_idx]
                _set_cell_margins(cell)
                cell.text = ""
                p = cell.paragraphs[0]
                run = p.add_run(value)
                run.font.size = Pt(10)

        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(8)

    def save(self, path: Path) -> None:
        self.doc.save(str(path))


def main() -> int:
    ASSETS.mkdir(parents=True, exist_ok=True)
    flujo = ASSETS / "flujo_clinico.png"
    draw_flujo_clinico(flujo)

    shots = ensure_screenshots()
    if not shots:
        print("Error: sin capturas. Sirva frontend (:8765) y backend (:8000).", file=sys.stderr)
        return 1

    print(f"Usando {len(shots)} capturas reales.")
    docx = InstructivoDOCX()
    render_instructivo(docx, shots, flujo)
    docx.save(OUTPUT)
    print(f"DOCX generado: {OUTPUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
